"""Step 1 of ingestion: turn an uploaded file into (page_number, text) pairs.

Library choices
- PDF  -> PyMuPDF (`import fitz`): fast, handles multi-column layouts well
          with sort=True, and gives true per-page text for citations.
- DOCX -> python-docx: paragraphs + tables. Word has no fixed pages, so we
          group text into ~3000-character "sections" and cite those.
- TXT/MD -> plain UTF-8 read, grouped the same way.
"""
import re
from dataclasses import dataclass

import docx
import fitz  # PyMuPDF

SECTION_CHARS = 3000


@dataclass
class LoadedPage:
    page_number: int
    text: str


def load_document(file_path: str, file_type: str) -> list[LoadedPage]:
    file_type = file_type.lower()
    if file_type == "pdf":
        pages = _load_pdf(file_path)
    elif file_type == "docx":
        pages = _load_docx(file_path)
    elif file_type in ("txt", "md"):
        pages = _group_into_sections(_read_text(file_path))
    else:
        raise ValueError(f"Unsupported file type: .{file_type}")

    pages = [LoadedPage(p.page_number, _clean(p.text)) for p in pages]
    pages = [p for p in pages if p.text]
    if not pages:
        raise ValueError(
            "No extractable text found. If this is a scanned PDF, it needs OCR "
            "(not enabled yet)."
        )
    return pages


def _load_pdf(path: str) -> list[LoadedPage]:
    pages = []
    with fitz.open(path) as pdf:
        for index, page in enumerate(pdf, start=1):
            pages.append(LoadedPage(index, page.get_text("text", sort=True)))
    return pages


def _load_docx(path: str) -> list[LoadedPage]:
    document = docx.Document(path)
    blocks = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return _group_into_sections("\n\n".join(blocks))


def _read_text(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def _group_into_sections(text: str) -> list[LoadedPage]:
    """Pack paragraphs into ~SECTION_CHARS blocks so citations point at a
    stable 'page' even for formats without pages."""
    sections, current, number = [], "", 1
    for paragraph in re.split(r"\n\s*\n", text):
        if current and len(current) + len(paragraph) > SECTION_CHARS:
            sections.append(LoadedPage(number, current))
            number += 1
            current = ""
        current = f"{current}\n\n{paragraph}".strip()
    if current.strip():
        sections.append(LoadedPage(number, current))
    return sections


def _clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"-\n(?=[a-z])", "", text)          # re-join hyphenated line breaks
    text = re.sub(r"[ \t]+", " ", text)                # collapse spaces
    text = re.sub(r"\n{3,}", "\n\n", text)             # collapse blank runs
    return text.strip()
